#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Word文档批量处理工具（支持递归处理子文件夹）
保存为 .py 或 .pyw 文件均可
"""

import os
import sys
import shutil
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from docx import Document
from docx.shared import Cm

# 如果是 Windows 系统，隐藏控制台窗口
if sys.platform == 'win32':
    try:
        import ctypes
        ctypes.windll.user32.ShowWindow(ctypes.windll.kernel32.GetConsoleWindow(), 0)
    except Exception:
        pass


class DocxProcessorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Word文档批量处理工具（含子文件夹）")
        self.root.geometry("700x660")
        self.root.resizable(False, False)
        
        try:
            self.root.iconbitmap(default='icon.ico')
        except Exception:
            pass
        
        self.root.configure(bg='#f0f0f0')
        
        # 主框架
        main_frame = tk.Frame(root, bg='#f0f0f0', padx=20, pady=20)
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # 标题
        title_label = tk.Label(
            main_frame, text="Word文档批量处理工具", 
            font=("微软雅黑", 18, "bold"), 
            bg='#f0f0f0', fg='#2c3e50'
        )
        title_label.pack(pady=(0, 15))
        
        # 文件夹选择区域
        folder_frame = tk.Frame(main_frame, bg='#f0f0f0')
        folder_frame.pack(fill=tk.X, pady=10)
        
        tk.Label(
            folder_frame, text="选择文件夹:", 
            font=("微软雅黑", 11), 
            bg='#f0f0f0', fg='#34495e'
        ).pack(side=tk.LEFT, padx=(0, 10))
        
        self.folder_path = tk.StringVar()
        self.folder_entry = tk.Entry(
            folder_frame, textvariable=self.folder_path, 
            width=45, font=("微软雅黑", 10),
            bg='#ffffff', fg='#2c3e50',
            relief=tk.GROOVE
        )
        self.folder_entry.pack(side=tk.LEFT, padx=(0, 10))
        
        self.browse_btn = tk.Button(
            folder_frame, text="浏览...", 
            command=self.browse_folder,
            bg='#3498db', fg='white', 
            font=("微软雅黑", 10, "bold"), 
            padx=15, pady=3,
            cursor="hand2",
            relief=tk.RAISED,
            bd=2
        )
        self.browse_btn.pack(side=tk.LEFT)
        
        # 处理设置说明
        info_frame = tk.LabelFrame(
            main_frame, text="处理设置", 
            font=("微软雅黑", 11, "bold"),
            bg='#ffffff', fg='#2c3e50',
            padx=15, pady=10,
            relief=tk.GROOVE
        )
        info_frame.pack(fill=tk.X, pady=10)
        
        settings_text = (
            "▸ 递归检索：包含选定文件夹及其所有子文件夹下的文档\n"
            "▸ 删除页眉：清空所有节的页眉内容\n"
            "▸ 页边距：上 1.27cm | 下 1.27cm | 左 1.27cm | 右 1.27cm\n"
            "▸ 页眉/页脚边距：页眉 0.8cm | 页脚 0.8cm\n"
            "▸ 备份机制：自动创建原层级关系的备份文件夹 (backup_before_modify)"
        )
        
        tk.Label(
            info_frame, text=settings_text, 
            font=("微软雅黑", 9), # 这里已经把 9.5 修改为整数 9，修复报错
            bg='#ffffff', fg='#34495e', 
            justify=tk.LEFT
        ).pack(anchor=tk.W)
        
        # 文件统计信息
        self.stats_frame = tk.Frame(main_frame, bg='#f0f0f0')
        self.stats_frame.pack(fill=tk.X, pady=5)
        
        self.file_count_label = tk.Label(
            self.stats_frame, text="⏳ 等待选择文件夹...", 
            font=("微软雅黑", 10),
            bg='#f0f0f0', fg='#7f8c8d'
        )
        self.file_count_label.pack(side=tk.LEFT)
        
        # 进度条区域
        self.progress_frame = tk.Frame(main_frame, bg='#f0f0f0')
        self.progress_frame.pack(fill=tk.X, pady=5)
        
        style = ttk.Style()
        style.theme_use('clam')
        style.configure(
            "green.Horizontal.TProgressbar", 
            background='#27ae60',
            troughcolor='#ecf0f1',
            bordercolor='#bdc3c7',
            lightcolor='#2ecc71',
            darkcolor='#27ae60'
        )
        
        self.progress_bar = ttk.Progressbar(
            self.progress_frame, 
            style="green.Horizontal.TProgressbar",
            mode='determinate', 
            length=660
        )
        self.progress_label = tk.Label(
            self.progress_frame, text="", 
            font=("微软雅黑", 9),
            bg='#f0f0f0', fg='#27ae60'
        )
        
        # 日志区域
        log_frame = tk.LabelFrame(
            main_frame, text="处理日志", 
            font=("微软雅黑", 11, "bold"),
            bg='#f0f0f0', fg='#2c3e50',
            padx=10, pady=10,
            relief=tk.GROOVE
        )
        log_frame.pack(fill=tk.BOTH, expand=True, pady=10)
        
        text_frame = tk.Frame(log_frame, bg='#ffffff')
        text_frame.pack(fill=tk.BOTH, expand=True)
        
        self.log_text = tk.Text(
            text_frame, height=8, 
            font=("Consolas", 9), 
            bg='#ffffff', fg='#2c3e50',
            wrap=tk.WORD,
            relief=tk.SUNKEN,
            bd=1
        )
        
        scrollbar = tk.Scrollbar(
            text_frame, command=self.log_text.yview,
            bg='#bdc3c7',
            activebackground='#95a5a6'
        )
        self.log_text.configure(yscrollcommand=scrollbar.set)
        
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        # 按钮区域
        button_frame = tk.Frame(main_frame, bg='#f0f0f0', height=50)
        button_frame.pack(fill=tk.X, pady=(5, 0))
        button_frame.pack_propagate(False)
        
        left_buttons = tk.Frame(button_frame, bg='#f0f0f0')
        left_buttons.pack(side=tk.LEFT)
        
        self.process_btn = tk.Button(
            left_buttons, text="开始处理", 
            command=self.start_processing,
            bg='#27ae60', fg='white', 
            font=("微软雅黑", 11, "bold"), 
            padx=25, pady=5,
            state=tk.DISABLED,
            cursor="hand2",
            relief=tk.RAISED,
            bd=2,
            activebackground='#2ecc71',
            activeforeground='white'
        )
        self.process_btn.pack(side=tk.LEFT, padx=(0, 10))
        
        self.clear_btn = tk.Button(
            left_buttons, text="清空日志", 
            command=self.clear_log,
            bg='#e67e22', fg='white', 
            font=("微软雅黑", 10), 
            padx=18, pady=5,
            cursor="hand2",
            relief=tk.RAISED,
            bd=2,
            activebackground='#f39c12',
            activeforeground='white'
        )
        self.clear_btn.pack(side=tk.LEFT)
        
        self.exit_btn = tk.Button(
            button_frame, text="退出", 
            command=self.quit_app,
            bg='#e74c3c', fg='white', 
            font=("微软雅黑", 10), 
            padx=25, pady=5,
            cursor="hand2",
            relief=tk.RAISED,
            bd=2,
            activebackground='#c0392b',
            activeforeground='white'
        )
        self.exit_btn.pack(side=tk.RIGHT)
        
        self.status_label = tk.Label(
            button_frame, text="● 就绪", 
            font=("微软雅黑", 10),
            bg='#f0f0f0', fg='#27ae60'
        )
        self.status_label.pack(side=tk.RIGHT, padx=15)
        
        # 底部状态栏
        self.status_bar = tk.Label(
            root, text="就绪 - 请选择文件夹", 
            bd=1, relief=tk.SUNKEN, 
            anchor=tk.W, padx=10, pady=2,
            bg='#34495e', fg='white',
            font=("微软雅黑", 9)
        )
        self.status_bar.pack(side=tk.BOTTOM, fill=tk.X)
        
        self.root.protocol("WM_DELETE_WINDOW", self.quit_app)

    def get_all_docx_files(self, root_folder):
        """递归获取所有docx文件，排除备份文件夹及Word临时文件"""
        docx_files = []
        backup_dir_name = "backup_before_modify"
        
        for dirpath, dirnames, filenames in os.walk(root_folder):
            if backup_dir_name in dirnames:
                dirnames.remove(backup_dir_name)
                
            for filename in filenames:
                if filename.endswith(".docx") and not filename.startswith("~$"):
                    full_path = os.path.join(dirpath, filename)
                    docx_files.append(full_path)
        return docx_files

    def browse_folder(self):
        """选择文件夹并检测文档总数"""
        folder = filedialog.askdirectory(title="选择包含 Word 文档的文件夹")
        if folder:
            self.folder_path.set(folder)
            self.process_btn.config(state=tk.NORMAL, bg='#27ae60')
            self.log(f"📁 已选择文件夹: {folder}")
            
            docx_files = self.get_all_docx_files(folder)
            file_count = len(docx_files)
            
            if file_count > 0:
                self.file_count_label.config(text=f"📄 共检索到 {file_count} 个 docx 文件（含子文件夹）", fg='#27ae60')
                self.log(f"✅ 共检索到 {file_count} 个 docx 文件（含子文件夹）")
                self.status_label.config(text=f"● 准备处理 {file_count} 个文件", fg='#27ae60')
                self.status_bar.config(text=f"准备处理 {file_count} 个文件 - 点击'开始处理'继续")
            else:
                self.file_count_label.config(text="⚠️ 警告：所选文件夹及子文件夹中均未找到 docx 文件", fg='#e67e22')
                self.log("⚠️ 警告：所选文件夹及子文件夹中均未找到 docx 文件")
                self.status_label.config(text="● 没有找到可处理的文件", fg='#e67e22')
                self.status_bar.config(text="没有找到 docx 文件 - 请重新选择文件夹")

    def log(self, message):
        """添加日志"""
        self.log_text.insert(tk.END, message + "\n")
        self.log_text.see(tk.END)
        self.root.update()

    def clear_log(self):
        """清空日志"""
        self.log_text.delete(1.0, tk.END)
        self.log("📋 日志已清空")

    def update_status(self, message, progress=None):
        """更新 UI 状态栏与进度"""
        self.status_bar.config(text=message)
        self.status_label.config(text=f"● {message}")
        if progress is not None:
            self.progress_bar['value'] = progress
            self.progress_label.config(text=f"{progress:.1f}%")
        self.root.update()

    def start_processing(self):
        """启动处理线程"""
        folder_path = self.folder_path.get()
        
        if not folder_path or not os.path.exists(folder_path):
            messagebox.showerror("错误", "请选择有效的文件夹路径")
            return
        
        docx_files = self.get_all_docx_files(folder_path)
        if not docx_files:
            messagebox.showwarning("警告", "没有找到待处理的 docx 文件")
            return
        
        result = messagebox.askyesno(
            "确认", 
            f"📊 共找到 {len(docx_files)} 个 docx 文件（含所有子文件夹）\n\n"
            "系统将在主目录下自动按原文件夹层级创建备份。\n"
            "是否立即开始处理？"
        )
        if not result:
            return
        
        self.process_btn.config(state=tk.DISABLED, bg='#95a5a6')
        self.browse_btn.config(state=tk.DISABLED, bg='#95a5a6')
        self.clear_btn.config(state=tk.DISABLED, bg='#95a5a6')
        
        self.progress_bar.pack(side=tk.LEFT, padx=(0, 10))
        self.progress_label.pack(side=tk.LEFT)
        self.progress_bar['value'] = 0
        self.progress_label.config(text="0%")
        
        self.clear_log()
        self.log("🚀 开始递归批量处理文档...")
        
        thread = threading.Thread(target=self.process_documents, args=(folder_path, docx_files))
        thread.daemon = True
        thread.start()

    def process_documents(self, root_folder, docx_files):
        """批量修改核心逻辑"""
        try:
            backup_root = os.path.join(root_folder, "backup_before_modify")
            if not os.path.exists(backup_root):
                os.makedirs(backup_root)
                self.log(f"📁 创建主备份根目录: {backup_root}")

            total_files = len(docx_files)
            processed = 0
            failed = []

            self.log("\n" + "═" * 50)

            for i, file_path in enumerate(docx_files):
                rel_path = os.path.relpath(file_path, root_folder)

                try:
                    progress = ((i + 1) / total_files) * 100
                    self.update_status(f"正在处理: {rel_path}", progress)
                    self.log(f"\n📄 处理 ({i+1}/{total_files}): {rel_path}")

                    # 1. 备份文件
                    backup_file_path = os.path.join(backup_root, rel_path)
                    backup_file_dir = os.path.dirname(backup_file_path)
                    if not os.path.exists(backup_file_dir):
                        os.makedirs(backup_file_dir)
                    
                    shutil.copy2(file_path, backup_file_path)
                    self.log("   ✅ 原文件已按层级备份")

                    # 2. 修改 Word 文档
                    doc = Document(file_path)

                    # 清空页眉
                    for section in doc.sections:
                        header = section.header
                        for paragraph in header.paragraphs:
                            paragraph.clear()

                    # 设置页边距
                    for section in doc.sections:
                        section.top_margin = Cm(1.27)
                        section.bottom_margin = Cm(1.27)
                        section.left_margin = Cm(1.27)
                        section.right_margin = Cm(1.27)
                        section.header_distance = Cm(0.8)
                        section.footer_distance = Cm(0.8)

                    # Save
                    doc.save(file_path)
                    self.log("   ✅ 页面格式与页眉修改完成")
                    processed += 1

                except Exception as e:
                    self.log(f"   ❌ 处理失败: {str(e)}")
                    failed.append(rel_path)

            self.update_status("处理完成", 100)
            self.log("\n" + "═" * 50)
            self.log("🎉 所有文档处理完成！")
            self.log(f"✅ 成功处理: {processed} 个文件")
            if failed:
                self.log(f"❌ 失败: {len(failed)} 个文件")
                for f in failed:
                    self.log(f"   • {f}")
            self.log(f"📁 备份文件保存在: {backup_root}")
            self.log("═" * 50)

            self.file_count_label.config(text=f"✅ 已处理 {processed} 个文件", fg='#27ae60')

            result_msg = f"处理完成！\n\n✅ 成功：{processed} 个文件"
            if failed:
                result_msg += f"\n❌ 失败：{len(failed)} 个文件"
            result_msg += f"\n\n📁 原层级备份已保存至：\n{backup_root}"

            messagebox.showinfo("处理完成", result_msg)

        except Exception as e:
            self.log(f"❌ 处理过程中出现严重错误: {str(e)}")
            messagebox.showerror("错误", f"处理中断：{str(e)}")

        finally:
            self.reset_ui()

    def reset_ui(self):
        self.process_btn.config(state=tk.NORMAL, bg='#27ae60')
        self.browse_btn.config(state=tk.NORMAL, bg='#3498db')
        self.clear_btn.config(state=tk.NORMAL, bg='#e67e22')
        self.update_status("就绪")
        self.root.after(2000, self.hide_progress_bar)

    def hide_progress_bar(self):
        self.progress_bar.pack_forget()
        self.progress_label.pack_forget()
        self.progress_bar['value'] = 0

    def quit_app(self):
        if messagebox.askokcancel("退出", "确定要退出程序吗？"):
            self.root.quit()
            self.root.destroy()


def main():
    root = tk.Tk()
    app = DocxProcessorApp(root)
    
    root.update_idletasks()
    width = root.winfo_width()
    height = root.winfo_height()
    x = (root.winfo_screenwidth() // 2) - (width // 2)
    y = (root.winfo_screenheight() // 2) - (height // 2)
    root.geometry(f'{width}x{height}+{x}+{y}')
    
    root.mainloop()


if __name__ == "__main__":
    main()