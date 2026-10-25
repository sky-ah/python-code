import os
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from docx import Document
from docx.shared import Cm
import glob
import threading

class DocxProcessorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Word文档批量处理工具")
        self.root.geometry("700x750")  # 增大窗口尺寸
        self.root.resizable(False, False)
        
        # 设置样式
        self.root.configure(bg='#f0f0f0')
        
        # 创建主框架
        main_frame = tk.Frame(root, bg='#f0f0f0', padx=20, pady=20)
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # 标题
        title_label = tk.Label(main_frame, text="Word文档批量处理工具", 
                               font=("微软雅黑", 18, "bold"), 
                               bg='#f0f0f0', fg='#333333')
        title_label.pack(pady=(0, 20))
        
        # 文件夹选择区域
        folder_frame = tk.Frame(main_frame, bg='#f0f0f0')
        folder_frame.pack(fill=tk.X, pady=10)
        
        tk.Label(folder_frame, text="选择文件夹:", 
                font=("微软雅黑", 11), 
                bg='#f0f0f0').pack(side=tk.LEFT, padx=(0, 10))
        
        self.folder_path = tk.StringVar()
        self.folder_entry = tk.Entry(folder_frame, textvariable=self.folder_path, 
                                     width=45, font=("微软雅黑", 10))
        self.folder_entry.pack(side=tk.LEFT, padx=(0, 10))
        
        self.browse_btn = tk.Button(folder_frame, text="浏览...", 
                                    command=self.browse_folder,
                                    bg='#4CAF50', fg='white', 
                                    font=("微软雅黑", 10, "bold"), 
                                    padx=15, pady=3,
                                    cursor="hand2")
        self.browse_btn.pack(side=tk.LEFT)
        
        # 设置信息显示
        info_frame = tk.LabelFrame(main_frame, text="处理设置", 
                                   font=("微软雅黑", 11, "bold"),
                                   bg='#ffffff', fg='#333333',
                                   padx=10, pady=10)
        info_frame.pack(fill=tk.X, pady=15)
        
        settings_text = """
        • 删除所有页眉内容
        • 页边距设置为：上 1.27cm | 下 1.27cm | 左 1.27cm | 右 1.27cm
        • 页眉上边距：0.8cm
        • 页脚下边距：0.8cm
        • 自动创建备份文件夹
        """
        
        tk.Label(info_frame, text=settings_text, 
                font=("微软雅黑", 10), 
                bg='#ffffff', fg='#555555', 
                justify=tk.LEFT).pack(anchor=tk.W)
        
        # 文件统计信息
        self.stats_frame = tk.Frame(main_frame, bg='#f0f0f0')
        self.stats_frame.pack(fill=tk.X, pady=5)
        
        self.file_count_label = tk.Label(self.stats_frame, text="未选择文件夹", 
                                         font=("微软雅黑", 10),
                                         bg='#f0f0f0', fg='#666666')
        self.file_count_label.pack(side=tk.LEFT)
        
        # 进度条区域
        self.progress_frame = tk.Frame(main_frame, bg='#f0f0f0')
        self.progress_frame.pack(fill=tk.X, pady=10)
        
        self.progress_bar = ttk.Progressbar(self.progress_frame, mode='determinate', length=660)
        self.progress_label = tk.Label(self.progress_frame, text="", 
                                       font=("微软雅黑", 9),
                                       bg='#f0f0f0', fg='#666666')
        
        # 日志区域
        log_frame = tk.LabelFrame(main_frame, text="处理日志", 
                                  font=("微软雅黑", 11, "bold"),
                                  bg='#f0f0f0', fg='#333333',
                                  padx=10, pady=10)
        log_frame.pack(fill=tk.BOTH, expand=True, pady=10)
        
        # 创建带滚动条的文本框
        text_frame = tk.Frame(log_frame, bg='#f0f0f0')
        text_frame.pack(fill=tk.BOTH, expand=True)
        
        self.log_text = tk.Text(text_frame, height=8, 
                                font=("Consolas", 9), 
                                bg='#ffffff', fg='#333333',
                                wrap=tk.WORD)
        scrollbar = tk.Scrollbar(text_frame, command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=scrollbar.set)
        
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        # 按钮区域 - 放在底部明显位置
        button_frame = tk.Frame(main_frame, bg='#f0f0f0', height=50)
        button_frame.pack(fill=tk.X, pady=(10, 0))
        button_frame.pack_propagate(False)  # 防止框架收缩
        
        # 左侧按钮组
        left_buttons = tk.Frame(button_frame, bg='#f0f0f0')
        left_buttons.pack(side=tk.LEFT)
        
        self.process_btn = tk.Button(left_buttons, text="开始处理", 
                                     command=self.start_processing,
                                     bg='#2196F3', fg='white', 
                                     font=("微软雅黑", 12, "bold"), 
                                     padx=30, pady=8,
                                     state=tk.DISABLED,
                                     cursor="hand2")
        self.process_btn.pack(side=tk.LEFT, padx=(0, 10))
        
        self.clear_btn = tk.Button(left_buttons, text="清空日志", 
                                   command=self.clear_log,
                                   bg='#FF9800', fg='white', 
                                   font=("微软雅黑", 11), 
                                   padx=20, pady=8,
                                   cursor="hand2")
        self.clear_btn.pack(side=tk.LEFT)
        
        # 右侧提示标签
        self.status_label = tk.Label(button_frame, text="就绪", 
                                     font=("微软雅黑", 10),
                                     bg='#f0f0f0', fg='#4CAF50')
        self.status_label.pack(side=tk.RIGHT, padx=10)
        
        # 底部状态栏
        self.status_bar = tk.Label(root, text="就绪", 
                                   bd=1, relief=tk.SUNKEN, 
                                   anchor=tk.W, padx=10, pady=2,
                                   bg='#e0e0e0', fg='#333333',
                                   font=("微软雅黑", 9))
        self.status_bar.pack(side=tk.BOTTOM, fill=tk.X)
    
    def browse_folder(self):
        """浏览文件夹"""
        folder = filedialog.askdirectory(title="选择包含Word文档的文件夹")
        if folder:
            self.folder_path.set(folder)
            self.process_btn.config(state=tk.NORMAL, bg='#2196F3')
            self.log(f"已选择文件夹: {folder}")
            
            # 检查文件夹中的docx文件
            docx_files = glob.glob(os.path.join(folder, "*.docx"))
            file_count = len(docx_files)
            
            if file_count > 0:
                self.file_count_label.config(text=f"找到 {file_count} 个docx文件", fg='#4CAF50')
                self.log(f"找到 {file_count} 个docx文件")
                self.status_label.config(text=f"准备处理 {file_count} 个文件", fg='#2196F3')
            else:
                self.file_count_label.config(text="警告：文件夹中没有找到docx文件", fg='#FF9800')
                self.log("警告：文件夹中没有找到docx文件")
                self.status_label.config(text="没有找到可处理的文件", fg='#FF9800')
    
    def log(self, message):
        """添加日志信息"""
        self.log_text.insert(tk.END, message + "\n")
        self.log_text.see(tk.END)
        self.root.update()
    
    def clear_log(self):
        """清空日志"""
        self.log_text.delete(1.0, tk.END)
        self.log("日志已清空")
    
    def update_status(self, message, progress=None):
        """更新状态"""
        self.status_bar.config(text=message)
        self.status_label.config(text=message)
        if progress is not None:
            self.progress_bar['value'] = progress
            self.progress_label.config(text=f"{progress:.1f}%")
        self.root.update()
    
    def start_processing(self):
        """开始处理文档"""
        folder_path = self.folder_path.get()
        
        if not folder_path or not os.path.exists(folder_path):
            messagebox.showerror("错误", "请选择有效的文件夹路径")
            return
        
        # 确认对话框
        docx_files = glob.glob(os.path.join(folder_path, "*.docx"))
        if not docx_files:
            messagebox.showwarning("警告", "文件夹中没有找到docx文件")
            return
        
        result = messagebox.askyesno("确认", 
                                     f"找到 {len(docx_files)} 个docx文件\n"
                                     "处理前会自动创建备份\n"
                                     "是否继续？")
        if not result:
            return
        
        # 禁用按钮，防止重复点击
        self.process_btn.config(state=tk.DISABLED, bg='#cccccc')
        self.browse_btn.config(state=tk.DISABLED, bg='#cccccc')
        self.clear_btn.config(state=tk.DISABLED, bg='#cccccc')
        
        # 显示进度条
        self.progress_bar.pack(side=tk.LEFT, padx=(0, 10))
        self.progress_label.pack(side=tk.LEFT)
        self.progress_bar['value'] = 0
        self.progress_label.config(text="0%")
        
        # 清空旧日志
        self.clear_log()
        self.log("开始处理文档...")
        
        # 在新线程中处理，避免界面卡顿
        thread = threading.Thread(target=self.process_documents, args=(folder_path,))
        thread.daemon = True
        thread.start()
    
    def process_documents(self, folder_path):
        """处理文档的主函数"""
        try:
            docx_files = glob.glob(os.path.join(folder_path, "*.docx"))
            
            # 创建备份文件夹
            backup_folder = os.path.join(folder_path, "backup_before_modify")
            if not os.path.exists(backup_folder):
                os.makedirs(backup_folder)
                self.log(f"创建备份文件夹: {backup_folder}")
            
            total_files = len(docx_files)
            processed = 0
            failed = []
            
            self.log("\n" + "="*50)
            
            for i, file_path in enumerate(docx_files):
                file_name = os.path.basename(file_path)
                
                try:
                    # 更新进度
                    progress = ((i + 1) / total_files) * 100
                    self.update_status(f"正在处理: {file_name}", progress)
                    self.log(f"\n处理 ({i+1}/{total_files}): {file_name}")
                    
                    # 创建备份
                    backup_path = os.path.join(backup_folder, file_name)
                    with open(file_path, 'rb') as original, open(backup_path, 'wb') as backup:
                        backup.write(original.read())
                    self.log(f"  ✓ 已备份")
                    
                    # 打开文档
                    doc = Document(file_path)
                    
                    # 去除页眉
                    for section in doc.sections:
                        header = section.header
                        for paragraph in header.paragraphs:
                            paragraph.clear()
                    
                    # 修改页边距（上下左右均为1.27cm）
                    for section in doc.sections:
                        section.top_margin = Cm(1.27)
                        section.bottom_margin = Cm(1.27)
                        section.left_margin = Cm(1.27)
                        section.right_margin = Cm(1.27)
                        section.header_distance = Cm(0.8)
                        section.footer_distance = Cm(0.8)
                    
                    # 保存
                    doc.save(file_path)
                    self.log(f"  ✓ 格式修改完成")
                    processed += 1
                    
                except Exception as e:
                    self.log(f"  ✗ 错误: {str(e)}")
                    failed.append(file_name)
            
            # 完成
            self.update_status("处理完成", 100)
            self.log("\n" + "="*50)
            self.log("处理完成！")
            self.log(f"成功处理: {processed} 个文件")
            if failed:
                self.log(f"失败: {len(failed)} 个文件")
                for f in failed:
                    self.log(f"  - {f}")
            self.log(f"备份文件保存在: {backup_folder}")
            self.log("="*50)
            
            # 更新文件计数
            self.file_count_label.config(text=f"已处理 {processed} 个文件", fg='#4CAF50')
            
            # 显示结果消息框
            result_msg = f"处理完成！\n成功：{processed} 个文件"
            if failed:
                result_msg += f"\n失败：{len(failed)} 个文件"
            result_msg += f"\n\n备份保存在：\n{backup_folder}"
            
            messagebox.showinfo("处理完成", result_msg)
            
        except Exception as e:
            self.log(f"处理过程中出现错误: {str(e)}")
            messagebox.showerror("错误", f"处理过程中出现错误：{str(e)}")
        
        finally:
            self.reset_ui()
    
    def reset_ui(self):
        """重置UI状态"""
        self.process_btn.config(state=tk.NORMAL, bg='#2196F3')
        self.browse_btn.config(state=tk.NORMAL, bg='#4CAF50')
        self.clear_btn.config(state=tk.NORMAL, bg='#FF9800')
        self.update_status("就绪")
        
        # 延迟移除进度条
        self.root.after(2000, self.hide_progress_bar)
    
    def hide_progress_bar(self):
        """隐藏进度条"""
        self.progress_bar.pack_forget()
        self.progress_label.pack_forget()
        self.progress_bar['value'] = 0

def main():
    """主函数"""
    root = tk.Tk()
    app = DocxProcessorApp(root)
    
    # 设置窗口居中
    root.update_idletasks()
    width = root.winfo_width()
    height = root.winfo_height()
    x = (root.winfo_screenwidth() // 2) - (width // 2)
    y = (root.winfo_screenheight() // 2) - (height // 2)
    root.geometry(f'{width}x{height}+{x}+{y}')
    
    root.mainloop()

if __name__ == "__main__":
    main()